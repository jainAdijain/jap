"""Resume handling (T1, U7): store the file, extract text locally, parse via the model."""

import hashlib
import io
import re
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from database.models import Profile, SkillAlias
from llm.redact import Sensitive, strip_placeholders

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_BYTES = 10 * 1024 * 1024


class BadResume(ValueError):
    pass


def detect_mime(filename: str, data: bytes) -> str:
    """PDF or DOCX, checked by extension and by the file's first bytes."""
    name = filename.lower()
    if name.endswith(".pdf") and data.startswith(b"%PDF"):
        return PDF
    if name.endswith(".docx") and data.startswith(b"PK"):
        return DOCX
    raise BadResume("Upload a PDF or DOCX file")


def store(data_dir: Path, version_id: str, mime: str, data: bytes) -> tuple[Path, str]:
    folder = data_dir / "resumes"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{version_id}{'.pdf' if mime == PDF else '.docx'}"
    path.write_bytes(data)
    return path, hashlib.sha256(data).hexdigest()


def _pdf_links(reader: PdfReader) -> list[str]:
    """The targets of clickable links, such as a word "LinkedIn" that hides its URL."""
    out: list[str] = []
    for page in reader.pages:
        for annot in page.get("/Annots") or []:
            action = annot.get_object().get("/A")
            uri = action.get_object().get("/URI") if action is not None else None
            if uri:
                out.append(str(uri))
    return out


def _docx_links(doc: Any) -> list[str]:
    return [
        rel.target_ref
        for rel in doc.part.rels.values()
        if rel.is_external and rel.reltype.endswith("/hyperlink")
    ]


def extract_text(path: Path, mime: str) -> str:
    """The resume's text, then a Links line with the web links hidden behind words, so
    they are found (Profile offers them) and redacted like the links written out."""
    data = path.read_bytes()
    if mime == PDF:
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        hidden = _pdf_links(reader)
    else:
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        hidden = _docx_links(doc)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    shown = text.lower()
    extra = []
    for url in hidden:
        bare = re.sub(r"^https?://(www\.)?", "", url.strip()).rstrip("/")
        if url.lower().startswith(("http://", "https://")) and bare.lower() not in shown:
            extra.append(url.strip())
            shown += " " + bare.lower()
    return f"{text}\n\nLinks: {' '.join(extra)}" if extra else text


def sensitive_from_profile(profile: Profile | None) -> Sensitive:
    if profile is None:
        return Sensitive()
    return Sensitive(
        name=profile.full_name,
        emails=[profile.email] if profile.email else [],
        phones=[profile.phone] if profile.phone else [],
        addresses=[profile.address] if profile.address else [],
        links=[link["url"] for link in profile.links or [] if link.get("url")],
    )


async def skill_aliases(session: AsyncSession) -> dict[str, str]:
    rows = await session.execute(select(SkillAlias.alias, SkillAlias.canonical))
    return {alias: canonical for alias, canonical in rows.all()}


def normalize_skills(skills: list[str], aliases: dict[str, str]) -> list[str]:
    """Lowercase, map aliases to canonical names (postgresql -> postgres), de-duplicate."""
    out: list[str] = []
    for raw in skills:
        skill = " ".join(strip_placeholders(raw).lower().split())
        skill = aliases.get(skill, skill)
        if skill and skill not in out:
            out.append(skill)
    return out
